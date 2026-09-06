from datetime import date
from io import BytesIO

from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.services.reporting import get_daily_report, get_profit_report


bp = Blueprint("reports", __name__, url_prefix="/reports")


@bp.route("/daily")
def daily():
    today = date.today().isoformat()
    start_date = request.args.get("start_date") or today
    end_date = request.args.get("end_date") or start_date
    summary, distribution = get_daily_report(start_date, end_date)
    return render_template(
        "reports/daily.html",
        summary=summary,
        distribution=distribution,
        start_date=summary["start_date"],
        end_date=summary["end_date"],
    )


@bp.route("/profit")
def profit():
    today = date.today()
    start_date = request.args.get("start_date") or today.isoformat()
    end_date = request.args.get("end_date") or today.isoformat()
    report = get_profit_report(start_date, end_date)
    return render_template("reports/profit.html", report=report)


@bp.route("/profit/export")
def profit_export():
    today = date.today()
    start_date = request.args.get("start_date") or today.isoformat()
    end_date = request.args.get("end_date") or today.isoformat()
    report = get_profit_report(start_date, end_date)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Kar Analizi"

    headers = [
        "Ürün",
        "Adet",
        "İndirim Sonrası Satış (KDV Dahil)",
        "İndirim (KDV Dahil)",
        "İndirim Oranı",
        "Maliyet (KDV Dahil)",
        "Kâr / Maliyet Oranı",
        "Satış Kâr Marjı",
    ]
    header_fill = PatternFill("solid", fgColor="C9A84C")
    header_font = Font(bold=True, color="2A1010")

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, row in enumerate(report["products"], start=2):
        worksheet.cell(row=row_index, column=1, value=f"{row.name}{' / ' + row.variant if row.variant else ''}")
        worksheet.cell(row=row_index, column=2, value=row.quantity)
        worksheet.cell(row=row_index, column=3, value=float(row.net_revenue))
        worksheet.cell(row=row_index, column=4, value=float(row.discount_amount))
        worksheet.cell(row=row_index, column=5, value=float(row.discount_rate) / 100)
        worksheet.cell(row=row_index, column=6, value=float(row.cost))
        worksheet.cell(row=row_index, column=7, value=float(row.cost_profit_rate) / 100)
        worksheet.cell(row=row_index, column=8, value=float(row.sales_margin_rate) / 100)

    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 8
    for column in ["C", "D", "F"]:
        worksheet.column_dimensions[column].width = 18
    for column in ["E", "G", "H"]:
        worksheet.column_dimensions[column].width = 14

    for row in worksheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = '#,##0.00 "₺"'
    for row in worksheet.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00 "₺"'
    for row in worksheet.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.00%"
    for row in worksheet.iter_rows(min_row=2, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = "0.00%"

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    filename = f"kar-analizi-{report['start_date'].isoformat()}-{report['end_date'].isoformat()}.xlsx"
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
