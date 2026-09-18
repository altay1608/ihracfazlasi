from io import BytesIO

from flask import Blueprint, current_app, flash, jsonify, redirect, render_template, request, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import or_
from sqlalchemy.exc import IntegrityError
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models import (
    Category,
    InventoryCount,
    InventoryCountLine,
    Product,
    ProductBarcode,
    RetailMultiplier,
    SupplierInvoiceLine,
    Variant,
)
from app.services.barcodes import generate_code128_svg
from app.services.access_control import has_permission
from app.services.inventory_history import record_inventory_movement
from app.services.product_inventory import (
    compute_sale_price,
    compute_sale_price_gross_target,
    format_unit_barcode,
    get_available_barcodes,
    get_customer_sale_price,
    get_multiplier_by_id,
    get_next_product_code,
    round_customer_price,
    resolve_unit_barcode,
    sync_product_sale_price,
    sync_product_barcodes,
)
from app.services.reporting import attach_stock_thresholds
from app.services.reference_data import (
    ensure_reference_data,
    get_category_choices,
    get_default_retail_multiplier,
    get_preferred_product_multiplier,
    get_retail_multiplier_choices,
    get_variant_choices,
)
from app.utils import is_ajax_request, is_modal_request, quantize_amount
from .forms import ProductForm


bp = Blueprint("products", __name__, url_prefix="/products")


def get_product_delete_block_message(product):
    if product.sale_items.count() or product.return_items.count():
        return "Bu ürün işlem geçmişinde kullanıldığı için silinemez."

    approved_count_exists = (
        InventoryCountLine.query.join(InventoryCount)
        .filter(
            InventoryCountLine.product_id == product.id,
            InventoryCount.status == "approved",
        )
        .first()
        is not None
    )
    if approved_count_exists:
        return "Bu ürün onaylanmış stok sayımında bulunduğu için silinemez."
    return None


def prepare_product_for_delete(product):
    draft_count_lines = (
        InventoryCountLine.query.join(InventoryCount)
        .filter(
            InventoryCountLine.product_id == product.id,
            InventoryCount.status != "approved",
        )
        .all()
    )
    for line in draft_count_lines:
        db.session.delete(line)

    SupplierInvoiceLine.query.filter_by(product_id=product.id).update(
        {SupplierInvoiceLine.product_id: None},
        synchronize_session=False,
    )
    # Delete count scans before the product's unit barcodes on PostgreSQL.
    db.session.flush()


def build_product_label_entries(product):
    available_barcodes = get_available_barcodes(product)
    if getattr(product, "barcode_mode", "unit") == "shared":
        available_barcodes = available_barcodes[:1]
    barcode_values = [item.barcode for item in available_barcodes]
    if not barcode_values:
        barcode_values = [format_unit_barcode(product.product_code, 1)]

    return [
        {
            "product": product,
            "barcode_value": barcode_value,
            # Unit suffixes (-01, -02, ...) must be encoded, not only printed as text.
            "barcode_svg": generate_code128_svg(barcode_value, compact=True),
        }
        for barcode_value in barcode_values
    ]


def get_products_listing():
    ensure_reference_data()
    search = request.args.get("search", "").strip()
    category = request.args.get("category", "").strip()
    stock_filter = request.args.get("stock", "").strip()
    threshold = current_app.config["LOW_STOCK_THRESHOLD"]

    query = Product.query
    if search:
        like = f"%{search}%"
        query = query.filter(
            or_(
                Product.name.ilike(like),
                Product.product_code.ilike(like),
                Product.variant.ilike(like),
                Product.category.ilike(like),
                Product.product_barcodes.any(ProductBarcode.barcode.ilike(like)),
            )
        )
    if category:
        query = query.filter(Product.category == category)
    products = query.order_by(Product.updated_at.desc()).all()
    attach_stock_thresholds(products, threshold)
    if stock_filter == "low" and has_permission("alerts.access"):
        products = [product for product in products if product.stock_quantity <= product.effective_low_stock_threshold]
    elif stock_filter == "in" and has_permission("alerts.access"):
        products = [product for product in products if product.stock_quantity > product.effective_low_stock_threshold]
    elif stock_filter == "out":
        products = [product for product in products if product.stock_quantity <= 0]
    elif stock_filter != "all":
        products = [product for product in products if product.stock_quantity > 0]
    categories = [item.name for item in Category.query.order_by(Category.name.asc()).all()]
    return products, categories, threshold


def render_product_form(form, page_title, product=None, status_code=200):
    template_name = "products/_form_modal.html" if is_modal_request() else "products/form.html"
    return (
        render_template(
            template_name,
            form=form,
            page_title=page_title,
            product=product,
        ),
        status_code,
    )


def render_upload_modal(result=None, status_code=200):
    return (
        render_template("products/_upload_modal.html", result=result),
        status_code,
    )


def populate_product_form_defaults(form, product=None):
    default_multiplier = (
        get_default_retail_multiplier()
        if product is not None
        else get_preferred_product_multiplier()
    )

    if product is None and not (form.category.data or "").strip():
        category_values = {value for value, _label in form.category.choices}
        if "Tişört" in category_values:
            form.category.data = "Tişört"

    if not (form.product_code.data or "").strip():
        form.product_code.data = product.product_code if product else get_next_product_code()

    if not (form.retail_multiplier_id.data or "").strip():
        form.retail_multiplier_id.data = (
            str(product.retail_multiplier_id)
            if product and product.retail_multiplier_id
            else (str(default_multiplier.id) if default_multiplier else "")
        )

    purchase_price = form.purchase_price.data or (product.purchase_price if product else 0)
    multiplier = get_multiplier_by_id(form.retail_multiplier_id.data) or default_multiplier
    if multiplier:
        form.sale_price.data = round_customer_price(compute_sale_price_gross_target(purchase_price, multiplier.multiplier))


def get_barcode_span(product):
    available = [item for item in product.product_barcodes if item.status == "available"]
    sold = [item for item in product.product_barcodes if item.status == "sold"]
    all_units = sorted(available + sold, key=lambda item: item.sequence_no)
    if not all_units:
        return format_unit_barcode(product.product_code, 1)
    if len(all_units) == 1:
        return all_units[0].barcode
    return f"{all_units[0].barcode} / {all_units[-1].barcode}"


def prepare_product_payload(form, existing_product=None):
    multiplier = get_multiplier_by_id(form.retail_multiplier_id.data)
    if not multiplier:
        raise ValueError("Geçerli bir perakende çarpanı seçin.")

    purchase_price = form.purchase_price.data
    sale_price = compute_sale_price(purchase_price, multiplier.multiplier)
    return {
        "name": title_case_product_name(form.name.data),
        "category": form.category.data.strip(),
        "barcode": form.product_code.data.strip(),
        "product_code": form.product_code.data.strip(),
        "purchase_price": purchase_price,
        "sale_price": sale_price,
        "stock_quantity": form.stock_quantity.data,
        "critical_stock_level": (
            form.critical_stock_level.data
            if has_permission("alerts.access")
            else (existing_product.critical_stock_level if existing_product else None)
        ),
        "variant": (form.variant.data or "").strip() or None,
        "barcode_mode": (form.barcode_mode.data or "unit").strip() or "unit",
        "retail_multiplier_id": multiplier.id,
    }


def title_case_product_name(value):
    """Normalize product names so each word starts with an uppercase letter."""
    words = " ".join(str(value or "").strip().split()).split(" ")
    turkish_upper = {"i": "İ", "ı": "I"}
    normalized = []
    for word in words:
        if not word:
            continue
        first = turkish_upper.get(word[0], word[0].upper())
        normalized.append(first + word[1:].lower())
    return " ".join(normalized)


def render_bulk_multiplier_modal(product_ids):
    return render_template(
        "products/_bulk_multiplier_modal.html",
        product_ids=product_ids,
        multiplier_choices=get_retail_multiplier_choices(),
    )


def get_selected_product_variants(form):
    selected = list(dict.fromkeys(form.variants.data or []))
    if not selected:
        form.variants.errors.append("En az bir beden veya Varyant Yok seçeneğini işaretleyin.")
        return None
    if "__none__" in selected and len(selected) > 1:
        form.variants.errors.append("Varyant Yok seçeneği bedenlerle birlikte kullanılamaz.")
        return None
    return [None] if selected == ["__none__"] else selected


@bp.route("/")
def index():
    products, categories, threshold = get_products_listing()
    if request.args.get("partial") == "table":
        return render_template(
            "products/_table_section.html",
            products=products,
            categories=categories,
            threshold=threshold,
            get_barcode_span=get_barcode_span,
            get_customer_sale_price=get_customer_sale_price,
        )
    return render_template(
        "products/index.html",
        products=products,
        categories=categories,
        threshold=threshold,
        get_barcode_span=get_barcode_span,
        get_customer_sale_price=get_customer_sale_price,
    )


@bp.route("/add", methods=["GET", "POST"])
def add():
    form = ProductForm()
    if request.method == "POST":
        form.variant.data = ""
    populate_product_form_defaults(form)

    if request.method == "GET" and is_modal_request():
        html, _ = render_product_form(form, "Ürün Ekle")
        return html

    if form.validate_on_submit():
        selected_variants = get_selected_product_variants(form)
        product_code = form.product_code.data.strip()
        existing = Product.query.filter(
            or_(Product.product_code == product_code, Product.barcode == product_code)
        ).first()
        if existing:
            form.product_code.errors.append("Bu ürün kodu zaten kayıtlı.")
        elif selected_variants:
            try:
                payload = prepare_product_payload(form)
            except ValueError as exc:
                form.retail_multiplier_id.errors.append(str(exc))
            else:
                created_products = []
                next_product_code = product_code
                used_product_codes = {product_code}
                automatic_code = int(get_next_product_code()) if len(selected_variants) > 1 else None
                for selected_variant in selected_variants:
                    variant_payload = dict(payload)
                    variant_payload.update(
                        barcode=next_product_code,
                        product_code=next_product_code,
                        variant=selected_variant,
                    )
                    # Bedenli giyim ürünleri stok sayımı ve satışta birim bazında
                    # izlenir. Yanlışlıkla ortak barkod seçilse bile her adet için
                    # ayrı barkod ve etiket üretimini koru.
                    if selected_variant:
                        variant_payload["barcode_mode"] = "unit"
                    product = Product(**variant_payload)
                    db.session.add(product)
                    db.session.flush()
                    sync_product_barcodes(product, product.stock_quantity, product.barcode_mode)
                    record_inventory_movement(
                        product,
                        transaction_type="product_opening",
                        quantity_before=0,
                        quantity_after=product.stock_quantity,
                        source_type="product",
                        source_id=product.id,
                        source_reference=f"Ürün #{product.id}",
                    )
                    created_products.append(product)
                    if automatic_code is not None:
                        while str(automatic_code) in used_product_codes:
                            automatic_code += 1
                        next_product_code = str(automatic_code)
                        used_product_codes.add(next_product_code)
                        automatic_code += 1

                db.session.commit()
                created_product_ids = [product.id for product in created_products]
                print_url = url_for("products.bulk_labels", product_ids=created_product_ids)
                message = (
                    f"{len(created_products)} beden başarıyla stoğa eklendi. Etiketler hazırlandı."
                    if len(created_products) > 1
                    else "Ürün başarıyla eklendi. Etiketi hazırlandı."
                )
                if is_ajax_request():
                    return jsonify(
                        {
                            "success": True,
                            "message": message,
                            "refresh_target": "#products-table-section",
                            "print_url": print_url,
                        }
                    )
                flash(message, "success")
                return redirect(print_url)

    populate_product_form_defaults(form)
    if is_ajax_request():
        html, status_code = render_product_form(form, "Ürün Ekle", status_code=400 if form.errors else 200)
        return jsonify({"success": False, "html": html}), status_code
    return render_product_form(form, "Ürün Ekle")


@bp.route("/<int:product_id>/edit", methods=["GET", "POST"])
def edit(product_id):
    product = Product.query.get_or_404(product_id)
    form = ProductForm(obj=product)
    if product.variant and product.variant not in {value for value, _ in form.variant.choices}:
        form.variant.choices.append((product.variant, f"{product.variant} (mevcut eski beden)"))
    if request.method == "GET":
        form.product_code.data = product.product_code
        if product.retail_multiplier_id:
            form.retail_multiplier_id.data = str(product.retail_multiplier_id)
        populate_product_form_defaults(form, product=product)

    if request.method == "GET" and is_modal_request():
        html, _ = render_product_form(form, "Ürün Düzenle", product=product)
        return html

    if form.validate_on_submit():
        product_code = form.product_code.data.strip()
        existing = (
            Product.query.filter(
                or_(Product.product_code == product_code, Product.barcode == product_code),
                Product.id != product.id,
            )
            .first()
        )
        if existing:
            form.product_code.errors.append("Bu ürün kodu başka bir üründe kullanılıyor.")
        else:
            try:
                payload = prepare_product_payload(form, existing_product=product)
            except ValueError as exc:
                form.retail_multiplier_id.errors.append(str(exc))
            else:
                previous_stock = product.stock_quantity
                for key, value in payload.items():
                    setattr(product, key, value)
                try:
                    sync_product_barcodes(product, form.stock_quantity.data, product.barcode_mode)
                except ValueError as exc:
                    product.stock_quantity = previous_stock
                    form.stock_quantity.errors.append(str(exc))
                else:
                    transaction_type = "manual_in" if product.stock_quantity > previous_stock else "manual_out"
                    record_inventory_movement(
                        product,
                        transaction_type=transaction_type,
                        quantity_before=previous_stock,
                        quantity_after=product.stock_quantity,
                        source_type="product",
                        source_id=product.id,
                        source_reference=f"Ürün #{product.id}",
                        note="Ürün kartı üzerinden stok düzenleme",
                    )
                    db.session.commit()
                    if is_ajax_request():
                        return jsonify(
                            {
                                "success": True,
                                "message": "Ürün güncellendi.",
                                "refresh_target": "#products-table-section",
                            }
                        )
                    flash("Ürün güncellendi.", "success")
                    return redirect(url_for("products.index"))

    populate_product_form_defaults(form, product=product)
    if is_ajax_request():
        html, status_code = render_product_form(
            form,
            "Ürün Düzenle",
            product=product,
            status_code=400 if form.errors else 200,
        )
        return jsonify({"success": False, "html": html}), status_code
    return render_product_form(form, "Ürün Düzenle", product=product)


@bp.route("/<int:product_id>/delete", methods=["POST"])
def delete(product_id):
    product = Product.query.get_or_404(product_id)
    message = get_product_delete_block_message(product)
    if message:
        if is_ajax_request():
            return jsonify({"success": False, "message": message}), 400
        flash(message, "error")
        return redirect(url_for("products.index"))

    try:
        prepare_product_for_delete(product)
        record_inventory_movement(
            product,
            transaction_type="product_delete",
            quantity_before=product.stock_quantity,
            quantity_after=0,
            source_type="product",
            source_id=product.id,
            source_reference=f"Ürün #{product.id}",
            note="Ürün kartı silindi",
        )
        db.session.delete(product)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        message = "Ürün bağlı bir işlem kaydında kullanıldığı için silinemedi."
        if is_ajax_request():
            return jsonify({"success": False, "message": message}), 409
        flash(message, "error")
        return redirect(url_for("products.index"))

    if is_ajax_request():
        return jsonify(
            {
                "success": True,
                "message": "Ürün silindi.",
                "refresh_target": "#products-table-section",
            }
        )

    flash("Ürün silindi.", "success")
    return redirect(url_for("products.index"))


@bp.route("/bulk/delete", methods=["POST"])
def bulk_delete():
    product_ids = [int(value) for value in request.form.getlist("product_ids") if str(value).isdigit()]
    if not product_ids:
        return jsonify({"success": False, "message": "Silinecek ürün seçilmedi."}), 400

    products = Product.query.filter(Product.id.in_(product_ids)).all()
    blocked = [product.name for product in products if get_product_delete_block_message(product)]
    if blocked:
        return jsonify({"success": False, "message": f"İşlem geçmişi olan ürünler silinemez: {', '.join(blocked[:4])}"}), 400

    try:
        for product in products:
            prepare_product_for_delete(product)
            record_inventory_movement(
                product,
                transaction_type="product_delete",
                quantity_before=product.stock_quantity,
                quantity_after=0,
                source_type="product",
                source_id=product.id,
                source_reference=f"Ürün #{product.id}",
                note="Toplu ürün silme",
            )
            db.session.delete(product)
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify(
            {"success": False, "message": "Seçilen ürünlerden biri bağlı bir işlem kaydında kullanıldığı için silinemedi."}
        ), 409
    return jsonify(
        {
            "success": True,
            "message": f"{len(products)} ürün silindi.",
            "refresh_target": "#products-table-section",
        }
    )


@bp.route("/bulk/labels")
def bulk_labels():
    product_ids = [int(value) for value in request.args.getlist("product_ids") if str(value).isdigit()]
    products_by_id = {product.id: product for product in Product.query.filter(Product.id.in_(product_ids)).all()}
    products = [products_by_id[product_id] for product_id in product_ids if product_id in products_by_id]
    labels = []
    for product in products:
        labels.extend(build_product_label_entries(product))
    return render_template("products/bulk_labels.html", labels=labels)


@bp.route("/bulk/multiplier", methods=["GET", "POST"])
def bulk_multiplier():
    product_ids = [int(value) for value in request.values.getlist("product_ids") if str(value).isdigit()]
    products = Product.query.filter(Product.id.in_(product_ids)).all()
    if not products:
        return ("", 404)

    if request.method == "POST":
        multiplier_id = request.form.get("retail_multiplier_id", "").strip()
        multiplier = get_multiplier_by_id(multiplier_id)
        if not multiplier:
            return jsonify({"success": False, "message": "Geçerli bir perakende çarpanı seçin."}), 400

        for product in products:
            product.retail_multiplier_id = multiplier.id
            product.sale_price = compute_sale_price(product.purchase_price, multiplier.multiplier)
        db.session.commit()
        return jsonify(
            {
                "success": True,
                "message": f"{len(products)} ürünün çarpanı güncellendi.",
                "refresh_target": "#products-table-section",
            }
        )

    html = render_bulk_multiplier_modal(product_ids)
    return html if is_modal_request() else html


@bp.route("/<int:product_id>/stock", methods=["POST"])
def update_stock(product_id):
    product = Product.query.get_or_404(product_id)
    payload = request.get_json(silent=True) or request.form
    try:
        stock_quantity = int(payload.get("stock_quantity"))
    except (TypeError, ValueError):
        return jsonify({"success": False, "message": "Geçerli bir stok miktarı girin."}), 400

    if stock_quantity < 0:
        return jsonify({"success": False, "message": "Stok negatif olamaz."}), 400

    previous_stock = product.stock_quantity
    sync_product_barcodes(product, stock_quantity, product.barcode_mode)
    record_inventory_movement(
        product,
        transaction_type="manual_in" if product.stock_quantity > previous_stock else "manual_out",
        quantity_before=previous_stock,
        quantity_after=product.stock_quantity,
        source_type="product",
        source_id=product.id,
        source_reference=f"Ürün #{product.id}",
        note="Stok alanından manuel güncelleme",
    )
    db.session.commit()
    threshold = current_app.config["LOW_STOCK_THRESHOLD"]
    attach_stock_thresholds([product], threshold)
    return jsonify(
        {
            "success": True,
            "message": "Stok güncellendi.",
            "stock_quantity": product.stock_quantity,
            "is_low_stock": product.stock_quantity <= product.effective_low_stock_threshold,
            "effective_threshold": product.effective_low_stock_threshold,
            "critical_threshold": product.effective_critical_stock_threshold,
        }
    )


@bp.route("/barcode/<string:barcode>")
def get_by_barcode(barcode):
    raw_barcode = str(barcode or "").strip()
    barcode_record = resolve_unit_barcode(raw_barcode, status="available")

    # Tek ortak barkod modunda etiket, ürün kodunun kendisidir.
    if not barcode_record:
        shared_product = Product.query.filter(
            Product.barcode == raw_barcode,
            Product.barcode_mode == "shared",
        ).first()
        if shared_product and shared_product.stock_quantity > 0:
            barcode_record = (
                ProductBarcode.query.filter_by(product_id=shared_product.id, status="available")
                .order_by(ProductBarcode.sequence_no.asc())
                .first()
            )

    if not barcode_record:
        unavailable_barcode_record = resolve_unit_barcode(raw_barcode)
        if unavailable_barcode_record:
            return (
                jsonify(
                    {
                        "success": False,
                        "message": "Uygun stok miktarı yok, satış gerçekleştiremezsiniz.",
                    }
                ),
                409,
            )
        return (
            jsonify(
                {
                    "success": False,
                    "message": "Birim barkodu bulunamadı. Etiketteki kodu -01 uzantısı dahil eksiksiz okutun.",
                }
            ),
            404,
        )

    if getattr(barcode_record.product, "barcode_mode", "unit") == "shared" and barcode_record.product.stock_quantity <= 0:
        return jsonify(
            {
                "success": False,
                "message": "Uygun stok miktarı yok, satış gerçekleştiremezsiniz.",
            }
        ), 409

    product = barcode_record.product
    return jsonify(
        {
            "success": True,
            "product": {
                "id": product.id,
                "name": product.name,
                "category": product.category,
                "barcode": barcode_record.barcode,
                "barcode_id": barcode_record.id,
                "product_code": product.product_code,
                "sale_price": float(product.sale_price),
                "stock_quantity": product.stock_quantity,
                "variant": product.variant,
                "barcode_mode": product.barcode_mode,
            },
        }
    )


@bp.route("/<int:product_id>/label")
def label(product_id):
    product = Product.query.get_or_404(product_id)
    labels = build_product_label_entries(product)
    return render_template("products/label.html", product=product, labels=labels)


@bp.route("/template/download")
def download_template():
    ensure_reference_data()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Urunler"

    headers = [
        "Ürün Adı",
        "Kategori",
        "Alış Fiyatı (₺)",
        "Perakende Çarpanı",
        "Satış Fiyatı (₺, otomatik)",
        "Stok Miktarı (beden başına)",
        "Eklenecek Bedenler (virgülle)",
        "Ürün Bazlı KSS (opsiyonel)",
        "Barkod Tipi",
    ]
    example_row = ["Örnek Erkek Gömlek", "Gömlek", "250.00", "Standart 1.80x", "450.00", "10", "S, M, L", "", "unit"]

    header_fill = PatternFill("solid", fgColor="C9A84C")
    header_font = Font(bold=True, color="2A1010")
    example_font = Font(italic=True, color="7A7A7A")

    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        worksheet.column_dimensions[cell.column_letter].width = 24

    for index, value in enumerate(example_row, start=1):
        cell = worksheet.cell(row=2, column=index, value=value)
        cell.font = example_font
        cell.fill = PatternFill("solid", fgColor="E8E8E8")

    categories = [name for name, _ in get_category_choices()]
    variants = [name for name, _ in get_variant_choices(include_blank=False)]
    multipliers = [item.name for item in RetailMultiplier.query.order_by(RetailMultiplier.multiplier.asc()).all()]

    if categories:
        category_validation = DataValidation(type="list", formula1=f'"{",".join(categories)}"', allow_blank=False)
        worksheet.add_data_validation(category_validation)
        category_validation.add("B2:B500")

    if variants:
        variant_validation = DataValidation(type="list", formula1=f'"{",".join(variants)}"', allow_blank=True)
        worksheet.add_data_validation(variant_validation)
        variant_validation.add("F2:F500")

    if multipliers:
        multiplier_validation = DataValidation(type="list", formula1=f'"{",".join(multipliers)}"', allow_blank=False)
        worksheet.add_data_validation(multiplier_validation)
        multiplier_validation.add("D2:D500")

    barcode_mode_validation = DataValidation(type="list", formula1='"unit,shared"', allow_blank=False)
    worksheet.add_data_validation(barcode_mode_validation)
    barcode_mode_validation.add("I2:I500")

    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return send_excel_file(stream, "ihrac_fazlasi_giyim_urun_sablonu.xlsx")


@bp.route("/template/upload", methods=["GET", "POST"])
def upload_template():
    if request.method == "GET":
        html, _ = render_upload_modal()
        return html if is_modal_request() else render_template("products/upload.html")

    uploaded_file = request.files.get("file")
    if not uploaded_file or not uploaded_file.filename:
        result = {"success": 0, "skipped": 0, "errors": ["Lütfen bir .xlsx dosyası seçin."]}
        if is_ajax_request():
            html, status_code = render_upload_modal(result=result, status_code=400)
            return jsonify({"success": False, "html": html, "message": "Dosya seçilmedi."}), status_code
        flash("Lütfen bir .xlsx dosyası seçin.", "error")
        return redirect(url_for("products.index"))

    filename = secure_filename(uploaded_file.filename)
    if not filename.lower().endswith(".xlsx"):
        result = {"success": 0, "skipped": 0, "errors": ["Yalnızca .xlsx uzantılı dosyalar kabul edilir."]}
        if is_ajax_request():
            html, status_code = render_upload_modal(result=result, status_code=400)
            return jsonify({"success": False, "html": html, "message": "Geçersiz dosya formatı."}), status_code
        flash("Yalnızca .xlsx uzantılı dosyalar kabul edilir.", "error")
        return redirect(url_for("products.index"))

    try:
        result = process_template_upload(uploaded_file)
    except Exception:
        result = {"success": 0, "skipped": 0, "errors": ["Dosya okunamadı veya geçerli bir Excel şablonu değil."]}
        if is_ajax_request():
            html, status_code = render_upload_modal(result=result, status_code=400)
            return jsonify({"success": False, "html": html, "message": "Dosya okunamadı."}), status_code
        flash("Dosya okunamadı veya geçerli bir Excel şablonu değil.", "error")
        return redirect(url_for("products.index"))

    if is_ajax_request():
        html, _ = render_upload_modal(result=result)
        return jsonify(
            {
                "success": True,
                "message": f"{result['success']} ürün eklendi, {result['skipped']} satır atlandı.",
                "html": html,
                "keep_open": True,
                "refresh_target": "#products-table-section",
            }
        )

    flash(f"{result['success']} ürün eklendi, {result['skipped']} satır atlandı.", "success")
    return redirect(url_for("products.index"))


def process_template_upload(uploaded_file):
    ensure_reference_data()
    workbook = load_workbook(uploaded_file, data_only=True)
    worksheet = workbook.active

    result = {"success": 0, "skipped": 0, "errors": []}
    fallback_category = Category.query.filter_by(name="Diğer").first()
    category_map = {item.name.lower(): item.name for item in Category.query.all()}
    variant_map = {item.name.lower(): item.name for item in Variant.query.all()}
    multiplier_map = {item.name.lower(): item for item in RetailMultiplier.query.all()}
    default_multiplier = get_default_retail_multiplier()

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        values = [("" if value is None else str(value).strip()) for value in row[:9]]
        if not any(values):
            continue
        if values == ["Örnek Erkek Gömlek", "Gömlek", "250.00", "Standart 1.80x", "450.00", "10", "S, M, L", "", "unit"]:
            result["skipped"] += 1
            continue

        name, category_name, purchase_price, multiplier_name, _sale_price_display, stock_quantity, variants_value, critical_stock_level, barcode_mode = values
        if not name or purchase_price == "" or stock_quantity == "":
            result["skipped"] += 1
            result["errors"].append(f"Satır {row_index}: zorunlu alanlar eksik.")
            continue

        try:
            purchase_price = float(purchase_price)
            stock_quantity = int(float(stock_quantity))
            critical_stock_level = int(float(critical_stock_level)) if critical_stock_level else None
        except ValueError:
            result["skipped"] += 1
            result["errors"].append(f"Satır {row_index}: fiyat veya stok alanı sayısal değil.")
            continue

        if purchase_price < 0 or stock_quantity < 0:
            result["skipped"] += 1
            result["errors"].append(f"Satır {row_index}: negatif değer kullanılamaz.")
            continue

        matched_category = category_map.get(category_name.lower()) if category_name else None
        multiplier = multiplier_map.get(multiplier_name.lower()) if multiplier_name else default_multiplier
        if not multiplier:
            result["skipped"] += 1
            result["errors"].append(f"Satır {row_index}: perakende çarpanı bulunamadı.")
            continue

        selected_variants = []
        for raw_variant in variants_value.replace(";", ",").replace("\n", ",").split(","):
            raw_variant = raw_variant.strip()
            if not raw_variant:
                continue
            matched_variant = variant_map.get(raw_variant.lower())
            if not matched_variant:
                result["errors"].append(f"Satır {row_index}: geçersiz beden/varyant: {raw_variant}.")
                selected_variants = []
                break
            if matched_variant not in selected_variants:
                selected_variants.append(matched_variant)
        if variants_value and not selected_variants:
            result["skipped"] += 1
            continue
        if barcode_mode not in {"unit", "shared"}:
            result["skipped"] += 1
            result["errors"].append(f"Satır {row_index}: barkod tipi unit veya shared olmalıdır.")
            continue
        for selected_variant in selected_variants or [None]:
            candidate_code = get_next_product_code()
            product = Product(
                name=title_case_product_name(name),
                category=matched_category or (fallback_category.name if fallback_category else "Diğer"),
                barcode=candidate_code,
                product_code=candidate_code,
                purchase_price=purchase_price,
                sale_price=compute_sale_price(purchase_price, multiplier.multiplier),
                stock_quantity=stock_quantity,
                critical_stock_level=critical_stock_level,
                variant=selected_variant,
                barcode_mode=barcode_mode,
                retail_multiplier_id=multiplier.id,
            )
            db.session.add(product)
            db.session.flush()
            sync_product_barcodes(product, stock_quantity, product.barcode_mode)
            record_inventory_movement(
                product,
                transaction_type="import_opening",
                quantity_before=0,
                quantity_after=product.stock_quantity,
                source_type="product_import",
                source_id=product.id,
                source_reference="Excel ürün şablonu",
            )
            result["success"] += 1

    db.session.commit()
    return result


def send_excel_file(stream, filename):
    from flask import send_file

    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
